"""Gera docs/agent-rules.docx, a especificação do módulo AGENT.

Rodar depois de mudar o classificador: `py -3 scripts/update-agent-rules-docx.py`
(precisa de python-docx). O catálogo CCCI é lido de src/agent/config.ts.
"""
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "agent-rules.docx"
CONFIG = (ROOT / "src" / "agent" / "config.ts").read_text(encoding="utf-8")
CCCI_ROWS = [
    (m.group(1), m.group(2), m.group(3), re.sub(r"\s+", " ", m.group(4)))
    for m in re.finditer(
        r'id: "([^"]+)",\s*domain: "([^"]+)",\s*label: "([^"]+)",\s*url: [^\n]+\n\s*hint:\s*"((?:[^"\\]|\\.)*)"',
        CONFIG,
    )
]
BLUE = "2856A3"
PALE = "EAF1FC"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_text(cell, value, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(value)
    r.bold = bold
    r.font.size = Pt(8.5)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], header, True, "FFFFFF")
        shade(t.rows[0].cells[i], BLUE)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if len(t.rows) % 2 == 1:
                shade(cells[i], "F7F9FC")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return t

def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)

def page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2); section.right_margin = Cm(2)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(5)
for name, size, color in [("Title", 25, BLUE), ("Heading 1", 17, BLUE), ("Heading 2", 12, "3C4B64")]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)
header = section.header.paragraphs[0]
header.text = "ConsBOT · Agent Mode"
header.runs[0].font.size = Pt(8); header.runs[0].font.color.rgb = RGBColor(100, 110, 125)
page_field(section.footer.paragraphs[0])

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Agent Mode")
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Especificação funcional e técnica · versão 3 · classificador na Decisions API")
r.italic = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(70, 85, 105)
doc.add_paragraph()
box = doc.add_table(rows=1, cols=1)
shade(box.cell(0, 0), PALE)
set_cell_text(box.cell(0, 0), "Objetivo: oferecer ações contextuais úteis sem substituir respostas substantivas, sem alegar resultados não consultados e com fallback sempre seguro para full.", True)
doc.add_paragraph()
doc.add_heading("Princípios", level=1)
for item in [
    "O turno passa por dois blocos de perguntas fechadas na Decisions API (POST /v1/decisions, gpt-6-luna), que não escreve texto: responde predicados (probabilidade) e escolhas entre opções dadas.",
    "CLASSIFICADOR, antes da resposta: decide só a rota do turno. O envio espera apenas por ele.",
    "DIRECIONADOR, disparado logo depois: decide os pills de ação numa chamada própria, sem segurar a resposta. Os pills aparecem quando a resposta termina. Usa a pergunta, não o texto respondido.",
    "Necessidade de resposta e relevância dos pills são decisões independentes. Um pill nunca substitui a resposta.",
    "A confiança de cada ação é a probabilidade do modelo, não um número escrito por ele. Os limiares vêm da suíte real.",
    "O termo de busca é sempre um trecho da mensagem do usuário: não há como inventar grafia nem concatenar palavras da resposta.",
    "Respostas sem o modelo principal (saudação, agradecimento, despedida, fontes carregadas, corpus) usam textos de catálogo.",
    "No Clássico, corpus é proibido.",
    "Qualquer falha cai no classificador antigo (Responses API) e, se ele também falhar, em full.",
    "Regras determinísticas locais (periódicos, livros em PDF) são só rede de segurança: entram quando o Direcionador não responde. No log real elas geravam 35 pills indevidos como complemento.",
]:
    bullet(item)

doc.add_page_break()
doc.add_heading("1. Operação do classificador", level=1)
table(["Parâmetro", "Contrato"], [
    ("Motor", "decisions (padrão). VITE_AGENT_CLASSIFIER_ENGINE=responses força o caminho antigo."),
    ("Rota", "Frontend → Main-Server POST /api/decisions → OpenAI POST /v1/decisions"),
    ("Modelo", "gpt-6-luna (único aceito pela Decisions API)"),
    ("Entrada", "Mensagem developer: contexto do ConsBOT + prompt do admin. Mensagem user: pergunta anterior ≤ 500, resposta anterior ≤ 900, mensagem atual."),
    ("Blocos", "Classificador: forma, lista_fontes (+ trechos_brutos em Citações). Direcionador: as demais perguntas, só nos turnos que vão ao modelo principal ou ao corpus."),
    ("Timeout", "4.000 ms no navegador; 3,5 s no servidor; sem retry."),
    ("Cache", "prompt_cache_key agent-classifier-v1-<apresentação> e agent-director-v1-<apresentação>"),
    ("Custo", "≈ 800 tokens no Classificador e ≈ 6.800 no Direcionador, US$ 0,10 por milhão; saída não é cobrada."),
    ("Latência medida", "Classificador (o que segura a resposta): p50 0,26 s. Direcionador, fora do caminho da resposta: p50 0,29 s. Antes da migração: 2,9 s · 15,5 s."),
    ("Cota", "/api/decisions não consome cota nem sofre o override de modelo por app do /api/llm."),
], [3.6, 12.8])

doc.add_heading("Questionário (forma, lista_fontes e trechos_brutos são do Classificador; as demais, do Direcionador)", level=2)
table(["Pergunta", "Tipo", "Decide"], [
    ("forma", "choice (12)", "Forma do turno: conceito, continuação, colagem, pergunta, procedência, busca literal, meta, fontes, saudação, agradecimento, despedida, outro."),
    ("busca_termo", "predicate", "Busca/localização explícita de termo, fonte de citação, ou listagem de verbetes por título/autor/especialidade."),
    ("escopo_busca", "choice (3)", "ambos (padrão) · livros · verbetes."),
    ("campo_verbete", "choice (5)", "texto · titulo · autor · especialidade · definologia."),
    ("obra", "choice (16)", "Obra do catálogo nomeada pelo usuário, ou nenhum."),
    ("busca_conscienciograma", "predicate", "Busca explícita no Conscienciograma."),
    ("bibliografia_livro", "predicate", "Referência de livro (citar, BEE, nome de obra sozinho)."),
    ("bibliografia_verbete", "predicate", "Referência de verbetes."),
    ("lexicons", "predicate", "Consulta lexical de palavra da língua comum (não neologismo)."),
    ("lista_fontes", "predicate", "Listar as fontes carregadas na conversa."),
    ("bibliomancia", "predicate", "Ortopensata ao acaso ou Bibliomancia."),
    ("encyclossapiens", "choice (4)", "Seção kit · enciclopedia · pesquisa, ou nenhum."),
    ("destino_<área>", "8 × choice", "Um destino do catálogo CCCI por área (mídia, atividades, livros, publicações, verbetes, referência, PGC, institucional), ou nenhum. O melhor entre as áreas vira pill."),
    ("recurso", "choice (8)", "periodicos · livros_pdf · enciclopedia · quiz · flashcards · consgpt · conslm, ou nenhum."),
    ("termo", "choice (≤ 255)", "Candidato extraído da mensagem, ou nenhum."),
    ("trechos_brutos", "predicate", "Só em Citações: pedido explícito de trechos literais."),
], [3.8, 2.6, 10.0])

doc.add_heading("Limiares", level=2)
table(["Condição", "Valor", "Constante"], [
    ("Predicado vira pill", "≥ 0,60", "AGENT_DECISION_THRESHOLD"),
    ("Escolha com «nenhum» vira pill", "1 − P(nenhum) ≥ 0,60", "AGENT_DECISION_THRESHOLD"),
    ("Destino CCCI vira pill", "1 − P(nenhum) na área ≥ 0,75", "AGENT_DECISION_CATALOG_MIN"),
    ("Turno responde direto", "forma social ou lista_fontes ≥ 0,75", "AGENT_DECISION_DIRECT_MIN"),
    ("Turno vai a corpus (Citações)", "trechos_brutos ≥ 0,70", "—"),
    ("Obra preenchida", "P(obra) ≥ 0,50", "—"),
], [5.6, 4.6, 6.2])
doc.add_paragraph("Nas escolhas com «nenhum», a probabilidade de que ALGUMA opção se aplique é 1 − P(nenhum); qual opção é o argmax das demais. Numa escolha grande a massa se espalha, e o melhor destino pode ter 0,4 com 1 − P(nenhum) = 0,9.")

doc.add_heading("Modalidades de resposta", level=2)
table(["responseMode", "Quando", "Resposta"], [
    ("full", "Padrão: toda pergunta, inclusive sobre o próprio ConsBOT.", "Modelo principal; pills ao lado."),
    ("direct", "Saudação, agradecimento ou despedida inequívocos e sem ação; ou lista_fontes.", "Texto de catálogo; lista local das fontes."),
    ("corpus", "Só em Citações, com pedido explícito de trechos brutos.", "Recuperação semântica literal."),
], [2.8, 7.4, 6.2])

doc.add_page_break()
doc.add_heading("2. Regras aplicadas em código", level=1)
doc.add_paragraph("Precedências que antes dependiam de o modelo lembrar uma regra do prompt agora são aplicadas na interpretação das respostas (src/agent/planner/decisions.ts):")
for item in [
    "Busca de termo sem escopo objetivo gera livros e verbetes juntos.",
    "Obra nomeada restringe a busca aos livros, salvo pedido explícito de verbetes.",
    "Busca no Conscienciograma suprime as buscas em livros e verbetes.",
    "Bibliografia de livro vence a busca quando é a leitura mais provável («Citar o Léxico»).",
    "O recurso periodicos encobre os destinos revistas e publicacoes_ccci.",
    "Um destino específico do catálogo encobre o recurso genérico: simulado da PGC em vez do quiz, livros em outros idiomas em vez do Drive em português.",
    "Pedido de referência bibliográfica encobre o catálogo e a autoria de livros; busca no Conscienciograma encobre a página do livro inteiro.",
    "A Encyclossapiens encobre as listagens de verbetes do catálogo CCCI.",
    "REGRA GERAL — uma necessidade, uma família de pill: quando a busca de termo e um destino do catálogo disparam juntos, pedido explícito de busca fica com a busca; nos demais, vence o mais provável.",
    "A mesma regra entre ferramentas próprias e áreas do catálogo: Encyclossapiens encobre a área verbetes, bibliografia de livro encobre a área livros, o recurso periodicos encobre a área publicações.",
    "A raiz do ICGE nunca vira pill só para haver uma ação.",
    "Verbetoteca do ICGE nunca aparece junto da busca de verbetes.",
    "lista_fontes é o turno inteiro: nenhum outro pill.",
    "Ferramenta que exige termo sem candidato válido não vira pill. No máximo três pills por turno.",
]:
    bullet(item)

doc.add_heading("Candidatos a termo", level=2)
for item in [
    "Trechos entre aspas primeiro; depois n-gramas de 1 a 6 palavras da mensagem atual; por fim os da pergunta anterior (continuação elíptica).",
    "Um termo nunca atravessa pontuação de frase e nunca começa nem termina em artigo, preposição, verbo de comando ou palavra que só nomeia o pedido (termo, palavra, livros, verbetes).",
    "Nome de fonte ou lugar de busca (Conscienciograma, Enciclopédia, Verbetoteca, ICGE, CCCI…) nunca é termo, nem sozinho nem em termo composto, salvo entre aspas. Verbos de autoria e de relação (escreveu, autor, relacionados) também não abrem nem fecham termo.",
    "Colagens com mais de 60 palavras só usam as 25 primeiras e as 25 últimas: nunca um termo de dentro do bloco.",
    "Até 254 candidatos, mais «nenhum». A grafia do usuário é preservada; não há correção ortográfica.",
]:
    bullet(item)

doc.add_heading("Coerência da resposta com os pills", level=2)
for item in [
    "Respostas diretas vêm de catálogo: não existe texto do classificador que possa prometer busca ou afirmar resultado.",
    "Em full, o modelo principal recebe contexto confiável com serviço, escopo e estado «ainda não consultado».",
    "Ausência no RAG atual é formulada como «não localizado nas fontes consultadas nesta resposta», distinta de inexistência.",
]:
    bullet(item)

doc.add_page_break()
doc.add_heading("3. Catálogo de capacidades", level=1)
table(["Intent", "Política", "Pergunta que decide"], [
    ("search_book", "fulfills_explicit_action", "busca_termo + escopo_busca + obra + termo"),
    ("search_verbete", "fulfills_explicit_action", "busca_termo + escopo_busca + campo_verbete + termo"),
    ("search_conscienciograma", "fulfills_explicit_action", "busca_conscienciograma + termo"),
    ("bibliografia_livros", "fulfills_explicit_action", "bibliografia_livro + obra"),
    ("bibliografia_verbetes", "fulfills_explicit_action", "bibliografia_verbete"),
    ("consulta_lexicons", "fulfills_explicit_action", "lexicons + termo"),
    ("bibliomancia", "fulfills_explicit_action", "bibliomancia"),
    ("encyclossapiens", "complementary", "encyclossapiens"),
    ("catalogo_ccci", "complementary", "destino_ccci"),
    ("open_resource", "fulfills_explicit_action", "recurso (regras determinísticas só no fallback)"),
    ("list_sources", "local", "lista_fontes"),
], [4.1, 4.4, 7.9])
doc.add_paragraph("ConsGPT e ConsLM só aparecem quando o usuário os nomeia. LexiCons sempre abre em Cosmovisão.")

doc.add_heading("Contratos de deep link", level=2)
table(["Serviço", "Contrato"], [
    ("Livros", "index_search_book.html?q=<termo>&books=<book_code>"),
    ("Verbetes", "index_search_verb.html?q=<termo>&fields=<todos|titulo|verbetografo|especialidade|definologia> (texto → todos, autor → verbetografo)"),
    ("Conscienciograma", "index_search_ccg.html?q=<termo>"),
    ("Bibliografia de livros", "index_biblio_wv.html?sigla=<sigla>&style=<bee|simples>"),
    ("Bibliografia de verbetes", "index_biblio_verbete.html?verbetes=<títulos>&style=<bee|simples>"),
    ("Bibliomancia", "index_mancia.html?autostart=1"),
    ("LexiCons", "?q=<termo>&autostart=1; sem parâmetro de modo"),
], [4.6, 11.8])

doc.add_heading("Identificadores de obras", level=2)
doc.add_paragraph("Os dois destinos do Cons-IA aceitam conjuntos diferentes de obras e descartam em silêncio o que não conhecem (conferido no código das páginas em 2026-10-19). «—» = a obra não existe naquele destino.")
table(["Id", "Busca em livros", "Sigla da bibliografia", "Obra"], [
    ("TEAT", "TEAT", "T200", "200 Teáticas"),
    ("EXP", "EXP", "EXP", "700 Experimentos"),
    ("DAC", "DAC", "DAC", "Dicionário de Argumentos"),
    ("HSP", "HSP", "HSP", "Homo sapiens pacificus"),
    ("HSR", "HSR", "HSR", "Homo sapiens reurbanisatus"),
    ("LO", "LO", "LO2", "Léxico de Ortopensatas (2ª ed.)"),
    ("MDE", "MDE", "MDE", "Manual da Dupla Evolutiva"),
    ("MP", "MP", "MP", "Manual da Proéxis"),
    ("TNP", "TNP", "TNP", "Manual da Tenepes"),
    ("PROJ", "PROJ", "PROJ", "Projeciologia"),
    ("TC", "TC", "TC", "Temas da Conscienciologia"),
    ("PROJ1986", "PROJ1986", "—", "Projeciologia (1986)"),
    ("ZEFIRO", "ZEFIRO", "—", "Zéfiro"),
    ("MINI_ARLINDO", "—", "—", "Minitertúlia — Arlindo"),
    ("QUEST", "—", "—", "Questões Mini"),
    ("PC", "—", "PC", "Projeções da Consciência"),
    ("CCG", "—", "CCG", "Conscienciograma"),
    ("T100", "—", "T100", "100 Testes da Conscienciometria"),
    ("NE", "—", "NE", "Nossa Evolução"),
    ("MRC", "—", "MRC", "Manual de Redação da Conscienciologia"),
    ("MMT", "—", "MMT", "Manual dos Megapensenes Trivocabulares"),
    ("DNC", "—", "DNC", "Dicionário de Neologismos"),
    ("EC10", "—", "EC10", "Enciclopédia da Conscienciologia (10ª ed.)"),
], [2.6, 3.0, 3.6, 7.2])

doc.add_page_break()
doc.add_heading("4. Catálogo CCCI", level=1)
doc.add_paragraph("Gerado de CCCI_DESTINATIONS em src/agent/config.ts. A descrição é o texto que o classificador recebe como descrição da opção em destino_ccci: é o que separa cada destino do vizinho.")
table(["area", "Área", "Rótulo", "Quando usar"], CCCI_ROWS, [3.4, 2.4, 3.4, 7.2])
doc.add_paragraph("Páginas vetadas do ICGE (page_id 4006, 6051, 1385) nunca viram pill, mesmo reintroduzidas por variável de ambiente.")
doc.add_paragraph("Origem: mapa completo do icge.org.br em 2026-10-19 (573 páginas pela API do WordPress, 180 úteis). Ficaram de fora as 363 páginas individuais do Quiz, versões antigas ou duplicadas (4623, 4034, 5830, 14106…) e páginas vazias ou protegidas por senha.")

doc.add_page_break()
doc.add_heading("5. Telemetria e fallbacks", level=1)
for item in [
    "Um turnId comum relaciona decisão, impressão, clique e feedback.",
    "reason é um rastro determinístico: forma e probabilidade, busca e escopo, destinos escolhidos, termo e recusas (ex.: decisions · forma=busca_literal(0.95) · busca=1:livros · termo=tenepes).",
    "classifierResponse guarda as respostas brutas da Decisions API, visíveis no diagnóstico do turno.",
    "Cada ação registra confidence (probabilidade), posição, serviço, destino e parâmetros.",
    "O Main-Server registra por chamada: app, ação, número de perguntas, tokens e duração.",
]:
    bullet(item)
table(["Falha", "Comportamento"], [
    ("/api/decisions responde 404/405 (Main-Server antigo)", "Classificador antigo pelo resto da sessão."),
    ("OpenAI responde 403 / not enabled", "Classificador antigo pelo resto da sessão."),
    ("Timeout, rede ou outro erro HTTP", "Classificador antigo neste turno."),
    ("Recusa em uma pergunta", "As demais seguem valendo; a recusa entra no reason."),
    ("Classificador antigo também falha", "full, sem pills."),
    ("Triagem passa de 3 s", "Turno segue em full; pills entram quando chegarem."),
], [7.2, 9.2])
doc.add_paragraph("O classificador antigo envia app=consbot-agent ao /api/llm, para não herdar o override de modelo e esforço do ConsBOT.")

doc.add_page_break()
doc.add_heading("6. Verificação", level=1)
table(["Comando", "O que mede"], [
    ("npm test", "Unitários: questionário dentro dos limites da API, candidatos a termo, interpretação das respostas, precedências, fallback de modo."),
    ("npm run agent:suite", "Suíte real contra o Main-Server (Decisions)."),
    ("npm run agent:suite -- --engine responses", "Mesma suíte no classificador antigo, para comparação."),
    ("npm run agent:suite -- --ficha controle", "20 frases escritas depois da calibragem."),
    ("npm run agent:suite -- --verbose", "Rastro de decisão de todos os casos."),
    ("npm run agent:suite -- --ficha controle_icge", "10 frases dos destinos novos, escritas antes da calibragem."),
    ("npm run agent:suite -- --ficha controle3", "12 frases escritas depois das regras gerais, para medir se generalizam."),
    ("npm run agent:suite -- --timing after", "Pills decididos depois da resposta, com o texto respondido como entrada."),
], [6.4, 10.0])
doc.add_heading("Resultado da migração (2026-10-19, 3 execuções por caso)", level=2)
table(["Métrica", "Responses (antigo)", "Decisions (novo)"], [
    ("Casos estáveis nas fichas (87)", "64", "87"),
    ("Casos instáveis entre execuções", "11", "0"),
    ("Ações indevidas", "32 em 96", "0 em 96"),
    ("Perguntas engolidas por direct", "3", "0"),
    ("Ficha controle, primeira passada (20)", "19", "19"),
    ("Latência p50 / p90", "2,9 s / 15,5 s", "0,27 s / 0,40 s"),
], [6.4, 5.0, 5.0])
doc.add_heading("Pills depois da resposta (avaliado e não adotado, 2026-10-19)", level=2)
doc.add_paragraph("Testou-se decidir os pills depois da resposta, com o texto respondido como entrada. Sobre respostas reais do app, ficou em 100 de 107 casos, contra 107 decidindo com a triagem; a ficha controle empatou em 20/20. Quando a resposta já explica o termo ou lista as instituições, o modelo larga o pill que o usuário pediu. O código e a medição ficaram disponíveis (--timing after).")
doc.add_heading("Catálogo CCCI ampliado (2026-10-19)", level=2)
doc.add_paragraph("Regras gerais para os erros do controle cego: catálogo dividido em oito perguntas por área; nome de fonte nunca é termo; catálogo e busca disputam a mesma necessidade. Suíte com 130 casos, 3 execuções: 130 estáveis e 0 ações indevidas em 111 execuções. Controle novo (controle3, 12 frases escritas depois das regras): 11/12; a falha é limítrofe («Explique o que é um Colégio Invisível» ficou sem o destino).")
doc.add_paragraph("Versão anterior, de 27 para 63 destinos. Suíte com 130 casos: 127 estáveis, 0 instáveis, 0 ações indevidas em 111 execuções; ficha icge 13/13; ficha controle_icge 8/10 (falhas: verbetes relacionados a um tema e ranking de verbetógrafos acompanhado de busca indevida).")
doc.add_paragraph("As fichas serviram para calibrar o questionário; a ficha controle mostra que, fora da amostra, a precisão empata. O ganho real está em estabilidade, ausência de pills espúrios e latência. Limitação conhecida nos dois motores: grafia errada do termo não é corrigida.")

doc.core_properties.title = "Agent Mode"
doc.core_properties.subject = "Regras e contratos do classificador ConsBOT (Decisions API)"
doc.core_properties.author = "ConsBOT"
doc.save(OUTPUT)
print(OUTPUT)
